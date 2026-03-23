import streamlit as st
import pandas as pd
import re
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# 1. 定義專業稽核紅線
TAGS_FRIED = ["炸", "酥", "裹粉", "爆", "脆", "可樂餅", "◎"] 
TAGS_PROCESSED = ["丸", "排", "素羊", "素火腿", "素肉", "獅子頭", "豆包", "炸豆腐", "△", "★"]
TAGS_FREQUENT = ["豆芽", "銀芽", "芽菜"]
TAGS_SEASONING = ["沙茶", "咖哩", "腐乳", "三杯", "麻婆", "糖醋"]

def process_excel(uploaded_file):
    # 讀取 Excel
    wb = load_workbook(uploaded_file)
    output_logs = []
    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        df = pd.DataFrame(ws.values)
        
        fried_count = 0
        veg_count = 0
        used_seasoning = set()
        
        # 偵測日期列與天數
        date_row = None
        days_count = 0
        for i in range(min(15, df.shape[0])):
            if any(k in str(df.iloc[i, 2]) for k in ["日期", "Date"]):
                date_row = i
                break
        
        if date_row is not None:
            for col in range(2, 7):
                if "202" in str(df.iloc[date_row, col]): days_count += 1

            # 開始逐格稽核 (C欄到G欄，即 index 2 到 6)
            for col_idx in range(3, 8): # Openpyxl 是 1-based
                for row_idx in range(1, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    val = str(cell.value).strip() if cell.value else ""
                    if len(val) < 1 or val == "None": continue

                    issue = None
                    # A. 炸物累計
                    if any(f in val for f in TAGS_FRIED):
                        fried_count += 1
                        limit = 1 if days_count < 5 else 1
                        if fried_count > limit:
                            issue = f"🚩炸物超標({fried_count}次)"

                    # B. 加工品規格
                    if any(p in val for p in TAGS_PROCESSED):
                        if not re.search(r"(\d+[xX*×]\d+)|(\d+\s*[gG克])", val):
                            issue = "⚠️規格未標註"

                    # C. 豆芽頻率
                    if any(v in val for v in TAGS_FREQUENT):
                        veg_count += 1
                        if veg_count > 1:
                            issue = "❌食材重複(豆芽)"

                    # 如果有問題，標色並紀錄
                    if issue:
                        cell.fill = red_fill
                        cell.font = white_font
                        output_logs.append({
                            "分頁": sheet_name,
                            "項目": val,
                            "問題": issue
                        })

    # 儲存到記憶體供下載
    virtual_workbook = BytesIO()
    wb.save(virtual_workbook)
    virtual_workbook.seek(0)
    return virtual_workbook, output_logs

# --- Streamlit 介面 ---
st.set_page_config(page_title="康橋膳食自動標記系統", layout="wide")
st.title("🛡️ 康橋林口：膳食稽核與「紅字標記」系統")
st.info("上傳後，系統會自動在 Excel 裡將違規格子「塗紅」，妳可以直接下載發給廠商。")

f = st.file_uploader("📂 上傳新北食品 Excel 菜單", type=["xlsx"])

if f:
    processed_file, logs = process_excel(f)
    
    if logs:
        st.error(f"🚨 發現 {len(logs)} 處違規！已自動在 Excel 中標記紅色。")
        st.table(pd.DataFrame(logs))
        
        # 下載按鈕
        st.download_button(
            label="📥 下載「紅字標記版」菜單給廠商",
            data=processed_file,
            file_name="新北食品菜單_審核修改版.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        st.success("✅ 檢查完畢，未發現違規項目。")
